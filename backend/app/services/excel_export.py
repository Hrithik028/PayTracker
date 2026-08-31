from __future__ import annotations

from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.config import PROJECT_ROOT, AppSettings
from app.services.reconciliation import reconcile


NAVY = "0F2747"
BLUE = "2563EB"
WHITE = "FFFFFF"
PALE_BLUE = "DBEAFE"
PALE_RED = "FEE2E2"
CURRENCY = '"$"#,##0.00'
HOURS = "0.00"
DATE = "dd mmm yyyy"


def generate_workbook(periods: list, settings: AppSettings) -> tuple[Path, str]:
    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    filename = f"paytracker-export-{timestamp}.xlsx"
    output = settings.exports_dir / filename
    workbook = Workbook()
    workbook.remove(workbook.active)

    _summary_sheet(workbook, periods)
    _payslips_sheet(workbook, periods)
    _items_sheet(workbook, periods)
    _timesheets_sheet(workbook, periods)
    _reconciliation_sheet(workbook, periods)
    _monthly_sheet(workbook, periods)
    workbook.properties.creator = "PayTracker"
    workbook.properties.description = (
        f"Local PayTracker export generated {datetime.now(timezone.utc).isoformat()}"
    )
    workbook.save(output)
    return output, output.relative_to(PROJECT_ROOT).as_posix()


def _summary_sheet(workbook: Workbook, periods: list) -> None:
    sheet = workbook.create_sheet("Summary")
    exported = datetime.now()
    sheet.append(["PayTracker confirmed records"])
    sheet.append(["Exported", exported])
    sheet.append([])
    headers = [
        "Employer",
        "Period Start",
        "Period End",
        "Gross Pay",
        "Net Pay",
        "Tax",
        "Superannuation",
        "Paid Hours",
        "Screenshot Hours",
    ]
    sheet.append(headers)
    for period in periods:
        payslip = period.payslip
        sheet.append(
            [
                period.employer.name if period.employer else "",
                period.start_date,
                period.end_date,
                payslip.gross_pay if payslip else None,
                payslip.net_pay if payslip else None,
                payslip.tax_withheld if payslip else None,
                payslip.superannuation if payslip else None,
                payslip.total_paid_hours if payslip else None,
                sum((shift.total_worked_hours or 0 for shift in period.shifts), Decimal("0")),
            ]
        )
    total_row = sheet.max_row + 1
    sheet.cell(total_row, 1, "Totals")
    for col in range(4, 10):
        letter = get_column_letter(col)
        sheet.cell(total_row, col, f"=SUM({letter}5:{letter}{total_row - 1})")
    _format_sheet(sheet, 4, currency_columns={4, 5, 6, 7}, hours_columns={8, 9}, date_columns={2, 3})
    sheet["A1"].font = Font(size=16, bold=True, color=NAVY)
    sheet["B2"].number_format = "dd mmm yyyy hh:mm"


def _payslips_sheet(workbook: Workbook, periods: list) -> None:
    sheet = workbook.create_sheet("Payslips")
    headers = [
        "Record ID", "Employer", "Employee", "Period Start", "Period End",
        "Payment Date", "Gross Pay", "Net Pay", "Tax", "Superannuation",
        "Paid Hours", "Deductions", "Allowances", "Notes",
    ]
    sheet.append(headers)
    for period in periods:
        payslip = period.payslip
        if not payslip:
            continue
        sheet.append([
            period.id, period.employer.name if period.employer else "",
            payslip.employee_name, period.start_date, period.end_date,
            payslip.payment_date, payslip.gross_pay, payslip.net_pay,
            payslip.tax_withheld, payslip.superannuation, payslip.total_paid_hours,
            payslip.deductions, payslip.allowances, payslip.notes,
        ])
    _format_sheet(sheet, 1, currency_columns={7, 8, 9, 10, 12, 13}, hours_columns={11}, date_columns={4, 5, 6})


def _items_sheet(workbook: Workbook, periods: list) -> None:
    sheet = workbook.create_sheet("Pay Items")
    sheet.append(["Record ID", "Employer", "Category", "Description", "Hours", "Hourly Rate", "Amount", "Verified"])
    for period in periods:
        for item in period.payslip.items if period.payslip else []:
            sheet.append([
                period.id, period.employer.name if period.employer else "",
                item.category, item.description, item.hours, item.hourly_rate,
                item.amount, "Yes" if item.verified else "No",
            ])
    _format_sheet(sheet, 1, currency_columns={6, 7}, hours_columns={5})


def _timesheets_sheet(workbook: Workbook, periods: list) -> None:
    sheet = workbook.create_sheet("Timesheets")
    sheet.append(["Record ID", "Employer", "Shift Date", "Day", "Start", "Finish", "Unpaid Break (min)", "Worked Hours", "Category", "Verified"])
    for period in periods:
        for shift in period.shifts:
            sheet.append([
                period.id, period.employer.name if period.employer else "",
                shift.shift_date,
                shift.shift_date.strftime("%A") if shift.shift_date else "",
                shift.start_time, shift.finish_time,
                shift.unpaid_break_minutes, shift.total_worked_hours, shift.category,
                "Yes" if shift.verified else "No",
            ])
    _format_sheet(sheet, 1, hours_columns={8}, date_columns={3})
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, 5).number_format = "hh:mm"
        sheet.cell(row, 6).number_format = "hh:mm"


def _reconciliation_sheet(workbook: Workbook, periods: list) -> None:
    sheet = workbook.create_sheet("Reconciliation")
    sheet.append([
        "Record ID",
        "Employer",
        "Payslip Hours",
        "Screenshot Hours",
        "Hours Difference",
        "Hours State",
        "Calculated Amount",
        "Paid Items",
        "Amount Difference",
        "Amount State",
        "Estimated Pay From Shifts",
        "Paid Compared Earnings",
        "Paid Minus Estimate",
        "Estimate Direction",
        "Estimate Complete",
        "Note",
    ])
    for period in periods:
        result = reconcile(period)
        sheet.append([
            period.id, period.employer.name if period.employer else "",
            Decimal(result["payslip_hours"]), Decimal(result["screenshot_hours"]),
            Decimal(result["hours_difference"]), result["hours_state"],
            Decimal(result["calculated_amount"]), Decimal(result["paid_line_item_amount"]),
            Decimal(result["amount_difference"]), result["amount_state"],
            Decimal(result["estimated_pay_from_shifts"]),
            Decimal(result["paid_compared_earnings"]),
            Decimal(result["estimated_pay_difference"]),
            result["estimated_pay_direction"].replace("_", " ").title(),
            "Yes" if result["estimated_pay_complete"] else "No",
            result["label"],
        ])
        if "discrepancy" in {
            result["hours_state"],
            result["amount_state"],
            result["estimated_pay_state"],
        }:
            for cell in sheet[sheet.max_row]:
                cell.fill = PatternFill("solid", fgColor=PALE_RED)
    _format_sheet(
        sheet,
        1,
        currency_columns={7, 8, 9, 11, 12, 13},
        hours_columns={3, 4, 5},
    )


def _monthly_sheet(workbook: Workbook, periods: list) -> None:
    sheet = workbook.create_sheet("Monthly Summary")
    sheet.append(["Month", "Gross Pay", "Net Pay", "Tax", "Superannuation", "Confirmed Hours"])
    grouped: dict[str, list[Decimal]] = {}
    for period in periods:
        reference = period.end_date or period.start_date
        if not reference or not period.payslip:
            continue
        values = grouped.setdefault(reference.strftime("%Y-%m"), [Decimal("0")] * 5)
        values[0] += period.payslip.gross_pay or 0
        values[1] += period.payslip.net_pay or 0
        values[2] += period.payslip.tax_withheld or 0
        values[3] += period.payslip.superannuation or 0
        values[4] += sum((shift.total_worked_hours or 0 for shift in period.shifts), Decimal("0"))
    for month, values in sorted(grouped.items()):
        sheet.append([month, *values])
    _format_sheet(sheet, 1, currency_columns={2, 3, 4, 5}, hours_columns={6})


def _format_sheet(
    sheet,
    header_row: int,
    *,
    currency_columns: set[int] = set(),
    hours_columns: set[int] = set(),
    date_columns: set[int] = set(),
) -> None:
    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(sheet.max_column)}{sheet.max_row}"
    for cell in sheet[header_row]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(bold=True, color=WHITE)
        cell.alignment = Alignment(vertical="center")
    for row in sheet.iter_rows(min_row=header_row + 1):
        for cell in row:
            if cell.column in currency_columns:
                cell.number_format = CURRENCY
            elif cell.column in hours_columns:
                cell.number_format = HOURS
            elif cell.column in date_columns:
                cell.number_format = DATE
    for column in range(1, sheet.max_column + 1):
        values = [str(sheet.cell(row, column).value or "") for row in range(1, min(sheet.max_row, 200) + 1)]
        sheet.column_dimensions[get_column_letter(column)].width = min(max(max(map(len, values), default=8) + 2, 11), 36)
