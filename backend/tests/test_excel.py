from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from app.models import Employer, PayPeriod, Payslip, PayslipItem, Shift
from app.services.excel_export import generate_workbook


class TempSettings:
    def __init__(self, path: Path) -> None:
        self.exports_dir = path


def test_excel_generation_has_six_formatted_sheets(tmp_path: Path) -> None:
    employer = Employer(name="Example Employer")
    period = PayPeriod(
        employer=employer,
        start_date=date(2026, 7, 1),
        end_date=date(2026, 7, 14),
        status="confirmed",
    )
    period.payslip = Payslip(
        gross_pay=Decimal("600"),
        net_pay=Decimal("500"),
        tax_withheld=Decimal("100"),
        total_paid_hours=Decimal("20"),
    )
    period.payslip.items.append(
        PayslipItem(category="Ordinary", description="Ordinary", hours=Decimal("20"), hourly_rate=Decimal("30"), amount=Decimal("600"), verified=True)
    )
    period.shifts.append(Shift(shift_date=date(2026, 7, 1), total_worked_hours=Decimal("8"), verified=True))
    path, _ = generate_workbook([period], TempSettings(tmp_path))
    workbook = load_workbook(path, data_only=False)
    assert workbook.sheetnames == ["Summary", "Payslips", "Pay Items", "Timesheets", "Reconciliation", "Monthly Summary"]
    assert workbook["Payslips"].freeze_panes == "A2"
    assert workbook["Summary"]["D5"].number_format == '"$"#,##0.00'
    assert workbook["Timesheets"]["D2"].value == "Wednesday"
    assert workbook["Reconciliation"]["K1"].value == "Estimated Pay From Shifts"
    assert workbook["Reconciliation"]["K2"].value == Decimal("240.00")
    assert workbook["Reconciliation"]["M2"].value == Decimal("360.00")
    assert workbook["Reconciliation"]["K2"].number_format == '"$"#,##0.00'
